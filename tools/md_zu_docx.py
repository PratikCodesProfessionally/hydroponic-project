#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Wandelt die Projekt-Dokumentation (README.md) in eine Word-Datei um.

    python tools/md_zu_docx.py                        # README.md -> Dokumentation.docx
    python tools/md_zu_docx.py README.md ausgabe.docx

Unterstuetzt die in der README verwendeten Markdown-Elemente:
Ueberschriften, Absaetze, Tabellen, Codebloecke, Zitate, Listen,
**fett**, `Code` und [Links](...).
"""

import re
import sys
from datetime import date

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("python-docx fehlt. Installieren mit:\n    python -m pip install python-docx")
    sys.exit(1)


# ---- Gestaltung ---------------------------------------------------
BLAU      = RGBColor(0x2A, 0x78, 0xD6)     # Akzent, wie im Dashboard
TINTE     = RGBColor(0x0B, 0x0B, 0x0B)
GRAU      = RGBColor(0x52, 0x51, 0x4E)
CODE_FONT = "Consolas"
CODE_HINTERGRUND = "F2F1EE"
TABELLE_KOPF     = "EAF1FB"

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+?\*\*|\[[^\]]+\]\([^)]+\))")
LINK   = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def absatz_schattierung(par, farbe):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), farbe)
    par._p.get_or_add_pPr().append(shd)


def zellen_schattierung(zelle, farbe):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), farbe)
    zelle._tc.get_or_add_tcPr().append(shd)


def linker_balken(par):
    pbdr = OxmlElement("w:pBdr")
    links = OxmlElement("w:left")
    links.set(qn("w:val"), "single")
    links.set(qn("w:sz"), "18")
    links.set(qn("w:color"), "2A78D6")
    links.set(qn("w:space"), "12")
    pbdr.append(links)
    par._p.get_or_add_pPr().append(pbdr)


def inline_schreiben(par, text, basis_fett=False):
    """Setzt einen Markdown-Text mit `Code`, **fett** und [Links](...) in Runs um."""
    for teil in INLINE.split(text):
        if not teil:
            continue
        if teil.startswith("`") and teil.endswith("`"):
            run = par.add_run(teil[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = TINTE
            run.bold = basis_fett
        elif teil.startswith("**") and teil.endswith("**"):
            inline_schreiben(par, teil[2:-2], basis_fett=True)
        elif LINK.fullmatch(teil):
            text_teil, ziel = LINK.fullmatch(teil).groups()
            # Gedruckt sind Anker und Dateipfade nicht klickbar - nur den Text setzen.
            run = par.add_run(text_teil)
            run.bold = basis_fett
            if not ziel.startswith("#"):
                run.font.name = CODE_FONT
                run.font.size = Pt(9.5)
        else:
            run = par.add_run(teil)
            run.bold = basis_fett


def dokument_anlegen():
    doc = Document()

    for sektion in doc.sections:
        sektion.page_width = Cm(21.0)
        sektion.page_height = Cm(29.7)
        sektion.left_margin = sektion.right_margin = Cm(2.5)
        sektion.top_margin = Cm(2.2)
        sektion.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TINTE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, groesse, farbe, davor in (
        ("Heading 1", 20, BLAU, 0),
        ("Heading 2", 14, TINTE, 18),
        ("Heading 3", 11.5, GRAU, 10),
    ):
        stil = doc.styles[name]
        stil.font.name = "Calibri"
        stil.font.size = Pt(groesse)
        stil.font.color.rgb = farbe
        stil.font.bold = True
        stil.paragraph_format.space_before = Pt(davor)
        stil.paragraph_format.space_after = Pt(6)

    # Seitenzahl in der Fusszeile
    fusszeile = doc.sections[0].footer.paragraphs[0]
    fusszeile.alignment = WD_ALIGN_PARAGRAPH.CENTER
    feld = OxmlElement("w:fldSimple")
    feld.set(qn("w:instr"), "PAGE")
    fusszeile._p.append(feld)

    return doc


def codeblock(doc, zeilen):
    par = doc.add_paragraph()
    absatz_schattierung(par, CODE_HINTERGRUND)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.left_indent = Cm(0.25)
    for i, zeile in enumerate(zeilen):
        if i:
            par.add_run().add_break()
        run = par.add_run(zeile)
        run.font.name = CODE_FONT
        run.font.size = Pt(9)


def tabelle(doc, kopf, reihen):
    tab = doc.add_table(rows=1 + len(reihen), cols=len(kopf))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT

    for spalte, text in enumerate(kopf):
        zelle = tab.rows[0].cells[spalte]
        zellen_schattierung(zelle, TABELLE_KOPF)
        par = zelle.paragraphs[0]
        inline_schreiben(par, text, basis_fett=True)
        par.paragraph_format.space_after = Pt(2)

    for r, reihe in enumerate(reihen, start=1):
        for spalte in range(len(kopf)):
            text = reihe[spalte] if spalte < len(reihe) else ""
            par = tab.rows[r].cells[spalte].paragraphs[0]
            inline_schreiben(par, text)
            par.paragraph_format.space_after = Pt(2)

    for reihe in tab.rows:
        for zelle in reihe.cells:
            for par in zelle.paragraphs:
                for run in par.runs:
                    if run.font.size is None:
                        run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def zellen_aufteilen(zeile):
    return [teil.strip() for teil in zeile.strip().strip("|").split("|")]


def umwandeln(md_pfad, docx_pfad):
    with open(md_pfad, "r", encoding="utf-8") as datei:
        zeilen = datei.read().splitlines()

    doc = dokument_anlegen()
    i = 0
    titel_gesetzt = False

    while i < len(zeilen):
        zeile = zeilen[i]

        # ---- Codeblock ----
        if zeile.startswith("```"):
            block = []
            i += 1
            while i < len(zeilen) and not zeilen[i].startswith("```"):
                block.append(zeilen[i])
                i += 1
            codeblock(doc, block)
            i += 1
            continue

        # ---- Tabelle ----
        if zeile.startswith("|") and i + 1 < len(zeilen) and set(zeilen[i + 1].replace("|", "").strip()) <= {"-", " ", ":"}:
            kopf = zellen_aufteilen(zeile)
            i += 2
            reihen = []
            while i < len(zeilen) and zeilen[i].startswith("|"):
                reihen.append(zellen_aufteilen(zeilen[i]))
                i += 1
            tabelle(doc, kopf, reihen)
            continue

        # ---- Zitat ----
        if zeile.startswith(">"):
            teile = []
            while i < len(zeilen) and zeilen[i].startswith(">"):
                teile.append(zeilen[i].lstrip("> ").rstrip())
                i += 1
            par = doc.add_paragraph()
            linker_balken(par)
            par.paragraph_format.left_indent = Cm(0.4)
            inline_schreiben(par, " ".join(t for t in teile if t))
            continue

        # ---- Ueberschriften (inline-Formatierung wie **fett** aufloesen) ----
        if zeile.startswith("### "):
            inline_schreiben(doc.add_heading("", level=3), zeile[4:])
            i += 1
            continue
        if zeile.startswith("## "):
            inline_schreiben(doc.add_heading("", level=2), zeile[3:])
            i += 1
            continue
        if zeile.startswith("# "):
            inline_schreiben(doc.add_heading("", level=1), zeile[2:])
            if not titel_gesetzt:
                datum = doc.add_paragraph(f"Stand: {date.today().strftime('%d.%m.%Y')}")
                for run in datum.runs:
                    run.font.color.rgb = GRAU
                    run.font.size = Pt(9.5)
                titel_gesetzt = True
            i += 1
            continue

        # ---- Trennlinie ----
        if zeile.strip() == "---":
            i += 1
            continue

        # ---- Listen ----
        if zeile.startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            inline_schreiben(par, zeile[2:])
            i += 1
            continue
        nummer = re.match(r"^(\d+)\.\s+(.*)", zeile)
        if nummer:
            par = doc.add_paragraph(style="List Number")
            inline_schreiben(par, nummer.group(2))
            i += 1
            continue

        # ---- Leerzeile ----
        if not zeile.strip():
            i += 1
            continue

        # ---- Absatz (Folgezeilen zusammenziehen) ----
        teile = [zeile.rstrip()]
        i += 1
        while i < len(zeilen) and zeilen[i].strip() and not re.match(
                r"^(#{1,3} |```|\||>|- |\d+\. |---$)", zeilen[i]):
            teile.append(zeilen[i].strip())
            i += 1
        par = doc.add_paragraph()
        inline_schreiben(par, " ".join(teile))

    doc.save(docx_pfad)
    print(f"{docx_pfad} geschrieben.")


if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else "README.md"
    ziel = sys.argv[2] if len(sys.argv) > 2 else "Dokumentation.docx"
    umwandeln(md, ziel)
